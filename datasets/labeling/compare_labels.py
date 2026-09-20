"""Concordância entre os dois rotuladores do prs_v1 (#40) e lista do que discutir.

    python3 datasets/labeling/compare_labels.py            # relatório + disagreements.md
    python3 datasets/labeling/compare_labels.py --partial  # aceita planilha incompleta

Lê prs_v1_bardi.xlsx (precisa de openpyxl) e prs_v1_claude.jsonl. Grava prs_v1_bardi.jsonl (a
planilha em formato versionável) e disagreements.md. O rótulo FINAL (datasets/prs_v1_labels.jsonl)
só nasce depois da discussão: concordâncias entram direto, discordâncias entram com a decisão.

Kappa de Cohen ao lado da concordância bruta: com classes desbalanceadas (quase tudo "não"),
95% de concordância pode ser acaso."""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

HERE = Path(__file__).resolve().parent
FLAGS = [
    "touches_auth_security",
    "touches_data_schema",
    "breaking_api",
    "touches_infra_ci",
    "has_tests",
    "description_explains_why",
]
FIELDS = ["change_type", "risk", *FLAGS, "via"]


def read_sheet(partial: bool) -> dict[str, dict[str, Any]]:
    from openpyxl import load_workbook

    ws = load_workbook(HERE / "prs_v1_bardi.xlsx", data_only=True)["rotulos"]
    head = [c.value for c in ws[1]]
    out: dict[str, dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(head, row, strict=False))
        if not rec.get("id") or rec["id"] == "ex":
            continue
        filled = [rec.get(f) not in (None, "") for f in FIELDS]
        if not any(filled):
            continue
        if not all(filled):
            if partial:
                continue
            sys.exit(
                f"{rec['id']}: linha incompleta ({[f for f, ok in zip(FIELDS, filled, strict=True) if not ok]})"
            )
        out[rec["id"]] = _normalize(rec)
    return out


def _normalize(rec: dict[str, Any]) -> dict[str, Any]:
    lab: dict[str, Any] = {
        "id": rec["id"],
        "change_type": rec["change_type"],
        "risk": int(rec["risk"]),
        "via": rec["via"],
    }
    for f in FLAGS:
        lab[f] = str(rec[f]).strip().lower() == "sim"
    for f in ("duvida", "fora_do_diff"):
        lab[f] = str(rec.get(f) or "").strip().lower() == "sim"
    lab["nota"] = rec.get("nota") or ""
    return lab


def kappa(pairs: list[tuple[Any, Any]]) -> float:
    n = len(pairs)
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in ca) / (n * n)
    return 1.0 if pe == 1 else (po - pe) / (1 - pe)


def report(bardi: dict[str, Any], claude: dict[str, Any]) -> list[str]:
    ids = sorted(set(bardi) & set(claude))
    print(f"{len(ids)} PRs rotulados pelos dois\n")
    print(
        f"{'pergunta':<26}{'concordam':>10}{'kappa':>8}   discordâncias mais comuns (bardi→claude)"
    )
    for f in FIELDS:
        pairs = [(bardi[i][f], claude[i][f]) for i in ids]
        diff = Counter(f"{a}→{b}" for a, b in pairs if a != b)
        agree = sum(a == b for a, b in pairs)
        print(f"{f:<26}{agree:>6}/{len(ids):<3}{kappa(pairs):>8.2f}   {dict(diff.most_common(4))}")
    doubt_b = sum(bardi[i]["duvida"] for i in ids)
    doubt_c = sum(claude[i]["duvida"] for i in ids)
    print(
        f"\ndúvida marcada: bardi {doubt_b} · claude {doubt_c} · os dois {sum(bardi[i]['duvida'] and claude[i]['duvida'] for i in ids)}"
    )
    return ids


def write_disagreements(ids: list[str], bardi: dict[str, Any], claude: dict[str, Any]) -> int:
    data = {}
    for line in (HERE.parent / "prs_v1.jsonl").read_text(encoding="utf-8").splitlines():
        r = json.loads(line)
        data[r["id"]] = r
    lines = [
        "# Discordâncias de rotulagem — prs_v1",
        "",
        "Decidir o rótulo final de cada linha. `?` = o rotulador marcou dúvida.",
        "",
    ]
    total = 0
    for i in ids:
        diffs = [f for f in FIELDS if bardi[i][f] != claude[i][f]]
        if not diffs:
            continue
        total += len(diffs)
        pr = data[i]
        lines.append(f"## {i} · [{pr['slug']}]({pr['url']}/files) · {pr['title'][:80]}")
        for f in diffs:
            lines.append(
                f"- **{f}**: bardi `{bardi[i][f]}`{'?' if bardi[i]['duvida'] else ''} · claude `{claude[i][f]}`{'?' if claude[i]['duvida'] else ''} · final: ___"
            )
        for who, lab in (("bardi", bardi[i]), ("claude", claude[i])):
            if lab["nota"]:
                lines.append(f"  - nota {who}: {lab['nota']}")
        lines.append("")
    (HERE / "disagreements.md").write_text("\n".join(lines), encoding="utf-8")
    return total


def main() -> None:
    partial = "--partial" in sys.argv
    bardi = read_sheet(partial)
    claude = {
        x["id"]: x
        for x in map(
            json.loads, (HERE / "prs_v1_claude.jsonl").read_text(encoding="utf-8").splitlines()
        )
    }
    (HERE / "prs_v1_bardi.jsonl").write_text(
        "\n".join(json.dumps(bardi[i], ensure_ascii=False) for i in sorted(bardi)) + "\n",
        encoding="utf-8",
    )
    ids = report(bardi, claude)
    cells = write_disagreements(ids, bardi, claude)
    print(
        f"\n{cells} células em discordância de {len(ids) * len(FIELDS)} → datasets/labeling/disagreements.md"
    )


if __name__ == "__main__":
    main()
